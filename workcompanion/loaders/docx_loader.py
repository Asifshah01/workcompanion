"""DOCX loader (Word documents)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from workcompanion.config.logging_config import get_logger
from workcompanion.loaders.base import BaseLoader, LoadError, paginate_text
from workcompanion.loaders.tables import table_to_markdown
from workcompanion.schemas.documents import ParsedDocument
from workcompanion.utils.text_utils import clean_text

logger = get_logger(__name__)


class DocxLoader(BaseLoader):
    """Parse ``.docx`` files, preserving heading hierarchy and tables."""

    extensions = (".docx", ".doc")
    source_type = "docx"

    def load(self, path: str | Path, **kwargs: Any) -> ParsedDocument:
        target = Path(path)
        if not target.is_file():
            raise LoadError(f"DOCX not found: {target}")
        try:
            import docx
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise LoadError("python-docx is not installed.") from exc

        try:
            document_obj = docx.Document(str(target))
        except Exception as exc:
            raise LoadError(f"Could not open '{target.name}' as a Word document: {exc}") from exc

        document = ParsedDocument(
            document_name=target.name,
            source_path=str(target.resolve()),
            source_type="docx",
        )
        self._read_core_properties(document_obj, document)

        blocks: list[str] = []
        tables: list[str] = []
        for paragraph in document_obj.paragraphs:
            text = clean_text(paragraph.text or "")
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "") or ""
            if style.lower().startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                blocks.append(f"{'#' * min(int(level), 6)} {text}")
            elif style.lower() in {"title", "subtitle"}:
                blocks.append(f"# {text}")
            else:
                blocks.append(text)

        for table in document_obj.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            rendered = table_to_markdown(rows)
            if rendered:
                tables.append(rendered)

        combined = "\n\n".join(blocks)
        if tables:
            combined += "\n\n" + "\n\n".join(tables)

        document.pages = paginate_text(combined)
        document.metadata["table_count"] = len(tables)
        document = self.finalize(document)
        if not document.title:
            document.title = self.detect_first_heading(document.pages)
        logger.info("Parsed DOCX '%s': %d blocks, %d tables", target.name, len(blocks), len(tables))
        return document

    @staticmethod
    def _read_core_properties(docx_document: Any, document: ParsedDocument) -> None:
        try:
            props = docx_document.core_properties
        except Exception:  # pragma: no cover
            return
        title = clean_text(getattr(props, "title", "") or "")
        author = clean_text(getattr(props, "author", "") or "")
        subject = clean_text(getattr(props, "subject", "") or "")
        if title:
            document.title = title
        if author:
            document.author = author
        if subject:
            document.subject = subject
        keywords = clean_text(getattr(props, "keywords", "") or "")
        if keywords:
            document.tags = [tag.strip() for tag in keywords.split(",") if tag.strip()][:12]