"""PPTX loader (PowerPoint decks): one page per slide."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from workcompanion.config.logging_config import get_logger
from workcompanion.loaders.base import BaseLoader, LoadError
from workcompanion.loaders.tables import table_to_markdown
from workcompanion.schemas.documents import ParsedDocument, RawPage
from workcompanion.utils.text_utils import clean_text

logger = get_logger(__name__)


class PptxLoader(BaseLoader):
    """Parse ``.pptx`` decks into slide pages with titles preserved."""

    extensions = (".pptx", ".ppt")
    source_type = "pptx"

    def load(self, path: str | Path, **kwargs: Any) -> ParsedDocument:
        target = Path(path)
        if not target.is_file():
            raise LoadError(f"PPTX not found: {target}")
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise LoadError("python-pptx is not installed.") from exc

        try:
            presentation = Presentation(str(target))
        except Exception as exc:
            raise LoadError(f"Could not open '{target.name}' as a PowerPoint file: {exc}") from exc

        document = ParsedDocument(
            document_name=target.name,
            source_path=str(target.resolve()),
            source_type="pptx",
        )
        self._read_core_properties(presentation, document)

        pages: list[RawPage] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title, lines, tables = self._read_slide(slide)
            text_parts: list[str] = []
            if title:
                text_parts.append(f"## Slide {slide_number}: {title}")
            text_parts.extend(lines)
            page = RawPage(
                page_number=slide_number,
                text="\n".join(text_parts),
                tables=tables,
                heading_hint=title,
            )
            if slide.has_notes_slide:
                try:
                    notes = clean_text(slide.notes_slide.notes_text_frame.text or "")
                    if notes:
                        page.text += f"\n\n**Speaker notes:** {notes}"
                except Exception:  # pragma: no cover
                    pass
            pages.append(page)

        document.pages = pages
        document = self.finalize(document)
        if not document.title:
            document.title = next(
                (page.heading_hint for page in document.pages if page.heading_hint), None
            )
        logger.info("Parsed PPTX '%s': %d slides", target.name, document.page_count)
        return document

    @staticmethod
    def _read_slide(slide: Any) -> tuple[str | None, list[str], list[str]]:
        title: str | None = None
        lines: list[str] = []
        tables: list[str] = []

        for shape in slide.shapes:
            try:
                is_placeholder = shape.is_placeholder
            except Exception:  # pragma: no cover
                is_placeholder = False

            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                rendered = table_to_markdown(rows)
                if rendered:
                    tables.append(rendered)
                continue

            if not getattr(shape, "has_text_frame", False):
                continue
            text = clean_text(shape.text_frame.text or "")
            if not text:
                continue
            if is_placeholder and shape.placeholder_format.type in (1, 3):  # TITLE / CENTER_TITLE
                title = title or text
                continue
            lines.extend(part for part in text.splitlines() if part.strip())

        if title is None:
            # Fall back to the first line if the title placeholder was missing.
            for line in lines:
                if len(line) < 120:
                    title = line
                    lines.remove(line)
                    break

        return title, lines, tables

    @staticmethod
    def _read_core_properties(presentation: Any, document: ParsedDocument) -> None:
        try:
            props = presentation.core_properties
        except Exception:  # pragma: no cover
            return
        title = clean_text(getattr(props, "title", "") or "")
        author = clean_text(getattr(props, "author", "") or "")
        subject = clean_text(getattr(props, "subject", "") or "")
        if title:
            document.title = title
        if author:
            document.author = author
        if subject:
            document.subject = subject